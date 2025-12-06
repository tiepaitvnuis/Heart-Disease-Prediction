import pandas as pd
import numpy as np
import joblib
import os
from sklearn.metrics import accuracy_score, classification_report, confusion_matrix, roc_auc_score, roc_curve, precision_score, recall_score, f1_score
from sklearn.model_selection import cross_val_score
import matplotlib.pyplot as plt
import seaborn as sns
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
import warnings
warnings.filterwarnings('ignore')


class SoftVotingEnsemble:
    """
    Soft Voting Ensemble Class

    OUTPUT METRICS:
    ---------------
    Ensemble sẽ output đầy đủ các metrics sau:
    - Accuracy: Độ chính xác tổng thể
    - Precision: Độ chính xác dương tính (weighted average)
    - Recall: Độ nhạy (Sensitivity) (weighted average)
    - F1-Score: Trung bình điều hòa Precision và Recall (weighted average)
    - ROC-AUC: Area Under ROC Curve

    Ngoài ra còn có per-class metrics cho từng class.
    """
    def __init__(self, models_info):
        """
        Khởi tạo Soft Voting Ensemble
        
        Parameters:
        -----------
        models_info : list of dict
            Danh sách thông tin các model
            Format: [
                {
                    'name': 'Model Name',
                    'model_path': 'path/to/model.pkl',
                    'scaler_path': 'path/to/scaler.pkl',  # Optional - nếu không có thì không scale
                    'weight': 1.0  # Trọng số cho voting (optional)
                },
                ...
            ]
        """
        self.models_info = models_info
        self.models = []
        self.scalers = []
        self.model_names = []
        self.weights = []
        self.X_test = None
        self.y_test = None
        
    def load_models(self):
        """Load tất cả các model và scaler"""
        print("Đang load các model...")
        
        for info in self.models_info:
            try:
                # Load model
                model = joblib.load(info['model_path'])
                
                # Load scaler nếu có, nếu không thì để None
                scaler = None
                if 'scaler_path' in info and info['scaler_path'] is not None:
                    try:
                        scaler = joblib.load(info['scaler_path'])
                        print(f"✓ Đã load {info['name']} (có scaler)")
                    except:
                        print(f"✓ Đã load {info['name']} (không có scaler)")
                else:
                    print(f"✓ Đã load {info['name']} (không dùng scaler)")
                
                self.models.append(model)
                self.scalers.append(scaler)
                self.model_names.append(info['name'])
                self.weights.append(info.get('weight', 1.0))
                
            except Exception as e:
                print(f"✗ Lỗi khi load {info['name']}: {str(e)}")
        
        # Normalize weights
        total_weight = sum(self.weights)
        self.weights = [w / total_weight for w in self.weights]
        
        print(f"\nTổng số model đã load: {len(self.models)}")
        print(f"Trọng số: {dict(zip(self.model_names, self.weights))}")
        
    def load_test_data(self, X_test, y_test):
        """Load dữ liệu test"""
        self.X_test = X_test
        self.y_test = y_test
        print(f"\nĐã load dữ liệu test: {X_test.shape}")
        
    def predict_proba(self, X):
        """
        Dự đoán xác suất bằng soft voting
        
        Returns:
        --------
        averaged_proba : array
            Xác suất trung bình có trọng số từ tất cả model
        individual_probas : list of arrays
            Xác suất từ từng model riêng lẻ
        """
        individual_probas = []
        
        for i, (model, scaler) in enumerate(zip(self.models, self.scalers)):
            # Scale data nếu có scaler, không thì dùng data gốc
            if scaler is not None:
                X_scaled = scaler.transform(X)
            else:
                X_scaled = X
            
            # Dự đoán xác suất
            proba = model.predict_proba(X_scaled)
            individual_probas.append(proba)
        
        # Tính trung bình có trọng số
        weighted_probas = [proba * weight for proba, weight in zip(individual_probas, self.weights)]
        averaged_proba = np.sum(weighted_probas, axis=0)
        
        return averaged_proba, individual_probas
    
    def predict(self, X):
        """Dự đoán class bằng soft voting"""
        averaged_proba, _ = self.predict_proba(X)
        predictions = np.argmax(averaged_proba, axis=1)
        return predictions, averaged_proba
    
    def evaluate(self):
        """Đánh giá ensemble model"""
        print("\n" + "="*60)
        print("KẾT QUẢ ĐÁNH GIÁ SOFT VOTING ENSEMBLE")
        print("="*60)
        
        # Dự đoán từ ensemble
        y_pred, y_proba = self.predict(self.X_test)
        
        # Dự đoán từ từng model riêng lẻ
        individual_predictions = []
        individual_accuracies = []
        
        for i, (model, scaler, name) in enumerate(zip(self.models, self.scalers, self.model_names)):
            # Scale data nếu có scaler
            if scaler is not None:
                X_scaled = scaler.transform(self.X_test)
            else:
                X_scaled = self.X_test
            
            pred = model.predict(X_scaled)
            acc = accuracy_score(self.y_test, pred)
            individual_predictions.append(pred)
            individual_accuracies.append(acc)
        
        # Tính toán metrics cho ensemble
        test_accuracy = accuracy_score(self.y_test, y_pred)
        precision = precision_score(self.y_test, y_pred, average='weighted')
        recall = recall_score(self.y_test, y_pred, average='weighted')
        f1 = f1_score(self.y_test, y_pred, average='weighted')
        
        # Precision, Recall, F1 cho từng class
        precision_per_class = precision_score(self.y_test, y_pred, average=None)
        recall_per_class = recall_score(self.y_test, y_pred, average=None)
        f1_per_class = f1_score(self.y_test, y_pred, average=None)
        
        # Confusion Matrix
        cm = confusion_matrix(self.y_test, y_pred)
        
        # ROC AUC Score
        roc_auc = roc_auc_score(self.y_test, y_proba[:, 1])
        
        # In kết quả
        print(f"\n{'='*60}")
        print("SOFT VOTING ENSEMBLE:")
        print(f"{'='*60}")
        print(f"Độ chính xác: {test_accuracy:.4f} ({test_accuracy*100:.2f}%)")
        print(f"Precision (weighted): {precision:.4f}")
        print(f"Recall (weighted): {recall:.4f}")
        print(f"F1-Score (weighted): {f1:.4f}")
        print(f"ROC AUC Score: {roc_auc:.4f}")
        
        print(f"\n{'='*60}")
        print("SO SÁNH VỚI TỪNG MODEL:")
        print(f"{'='*60}")
        for name, acc in zip(self.model_names, individual_accuracies):
            print(f"{name:20s}: {acc:.4f} ({acc*100:.2f}%)")
        print(f"{'Ensemble':20s}: {test_accuracy:.4f} ({test_accuracy*100:.2f}%)")
        
        improvement = test_accuracy - max(individual_accuracies)
        print(f"\nCải thiện so với model tốt nhất: {improvement:+.4f} ({improvement*100:+.2f}%)")
        
        # Classification report
        print("\n" + "-"*60)
        print("CLASSIFICATION REPORT (Test Set):")
        print("-"*60)
        print(classification_report(self.y_test, y_pred, 
                                   target_names=['Không bệnh (0)', 'Có bệnh (1)']))
        
        print("\nConfusion Matrix:")
        print(cm)
        
        return {
            'test_accuracy': test_accuracy,
            'precision': precision,
            'recall': recall,
            'f1_score': f1,
            'precision_per_class': precision_per_class,
            'recall_per_class': recall_per_class,
            'f1_per_class': f1_per_class,
            'roc_auc': roc_auc,
            'confusion_matrix': cm,
            'individual_accuracies': dict(zip(self.model_names, individual_accuracies)),
            'y_pred': y_pred,
            'y_proba': y_proba
        }
    
    def plot_confusion_matrix(self, save_path='ensemble_confusion_matrix.png'):
        """Vẽ confusion matrix"""
        y_pred, _ = self.predict(self.X_test)
        cm = confusion_matrix(self.y_test, y_pred)
        
        plt.figure(figsize=(8, 6))
        sns.heatmap(cm, annot=True, fmt='d', cmap='Purples', 
                   xticklabels=['Không bệnh', 'Có bệnh'],
                   yticklabels=['Không bệnh', 'Có bệnh'])
        plt.title('Confusion Matrix - Soft Voting Ensemble')
        plt.ylabel('Actual')
        plt.xlabel('Predicted')
        plt.tight_layout()
        plt.savefig(save_path, dpi=300, bbox_inches='tight')
        print(f"\nĐã lưu confusion matrix tại: {save_path}")
        plt.close()
        
    def plot_roc_curve(self, save_path='ensemble_roc_curve.png'):
        """Vẽ ROC curve cho ensemble và từng model"""
        plt.figure(figsize=(10, 8))
        
        # ROC curve cho ensemble
        y_pred, y_proba = self.predict(self.X_test)
        fpr_ensemble, tpr_ensemble, _ = roc_curve(self.y_test, y_proba[:, 1])
        roc_auc_ensemble = roc_auc_score(self.y_test, y_proba[:, 1])
        
        plt.plot(fpr_ensemble, tpr_ensemble, color='purple', lw=3, 
                label=f'Ensemble (AUC = {roc_auc_ensemble:.4f})', linestyle='-')
        
        # ROC curve cho từng model
        colors = ['blue', 'green', 'orange', 'red', 'brown', 'pink']
        for i, (model, scaler, name) in enumerate(zip(self.models, self.scalers, self.model_names)):
            # Scale data nếu có scaler
            if scaler is not None:
                X_scaled = scaler.transform(self.X_test)
            else:
                X_scaled = self.X_test
            
            proba = model.predict_proba(X_scaled)[:, 1]
            fpr, tpr, _ = roc_curve(self.y_test, proba)
            roc_auc = roc_auc_score(self.y_test, proba)
            
            plt.plot(fpr, tpr, color=colors[i % len(colors)], lw=2, 
                    label=f'{name} (AUC = {roc_auc:.4f})', linestyle='--', alpha=0.7)
        
        # Đường random
        plt.plot([0, 1], [0, 1], color='navy', lw=2, linestyle=':', label='Random')
        
        plt.xlim([0.0, 1.0])
        plt.ylim([0.0, 1.05])
        plt.xlabel('False Positive Rate')
        plt.ylabel('True Positive Rate')
        plt.title('ROC Curves - Soft Voting Ensemble vs Individual Models')
        plt.legend(loc="lower right")
        plt.grid(alpha=0.3)
        plt.tight_layout()
        plt.savefig(save_path, dpi=300, bbox_inches='tight')
        print(f"Đã lưu ROC curve tại: {save_path}")
        plt.close()
    
    def plot_model_comparison(self, results, save_path='ensemble_model_comparison.png'):
        """Vẽ biểu đồ so sánh accuracy của các model"""
        accuracies = list(results['individual_accuracies'].values())
        accuracies.append(results['test_accuracy'])
        
        names = list(results['individual_accuracies'].keys())
        names.append('Ensemble')
        
        colors = ['skyblue'] * len(self.model_names) + ['purple']
        
        plt.figure(figsize=(10, 6))
        bars = plt.bar(range(len(names)), accuracies, color=colors, edgecolor='black', linewidth=1.5)
        
        # Thêm giá trị trên mỗi cột
        for i, (bar, acc) in enumerate(zip(bars, accuracies)):
            height = bar.get_height()
            plt.text(bar.get_x() + bar.get_width()/2., height,
                    f'{acc:.4f}\n({acc*100:.2f}%)',
                    ha='center', va='bottom', fontsize=10, fontweight='bold')
        
        plt.xticks(range(len(names)), names, rotation=45, ha='right')
        plt.ylabel('Accuracy')
        plt.title('Model Accuracy Comparison')
        plt.ylim([0, 1.0])
        plt.grid(axis='y', alpha=0.3)
        plt.tight_layout()
        plt.savefig(save_path, dpi=300, bbox_inches='tight')
        print(f"Đã lưu biểu đồ so sánh tại: {save_path}")
        plt.close()
    
    def export_metrics_to_excel(self, results, filename='ensemble_metrics.xlsx'):
        """Xuất các metrics ra file Excel"""
        print(f"\nĐang xuất kết quả ra file Excel...")
        
        wb = Workbook()
        ws = wb.active
        ws.title = "Ensemble Metrics"
        
        # Định dạng style
        header_fill = PatternFill(start_color="9966CC", end_color="9966CC", fill_type="solid")
        header_font = Font(bold=True, color="FFFFFF", size=12)
        cell_alignment = Alignment(horizontal="center", vertical="center")
        border = Border(
            left=Side(style='thin'),
            right=Side(style='thin'),
            top=Side(style='thin'),
            bottom=Side(style='thin')
        )
        
        # Tiêu đề chính
        ws.merge_cells('A1:B1')
        ws['A1'] = 'SOFT VOTING ENSEMBLE - EVALUATION METRICS'
        ws['A1'].font = Font(bold=True, size=14)
        ws['A1'].alignment = cell_alignment
        ws['A1'].fill = PatternFill(start_color="663399", end_color="663399", fill_type="solid")
        ws['A1'].font = Font(bold=True, color="FFFFFF", size=14)
        
        # Header cho metrics tổng quan
        ws['A3'] = 'Metric'
        ws['B3'] = 'Value'
        ws['A3'].fill = header_fill
        ws['B3'].fill = header_fill
        ws['A3'].font = header_font
        ws['B3'].font = header_font
        ws['A3'].alignment = cell_alignment
        ws['B3'].alignment = cell_alignment
        
        # Dữ liệu metrics
        metrics_data = [
            ['Accuracy', f"{results['test_accuracy']:.4f}"],
            ['Precision', f"{results['precision']:.4f}"],
            ['Recall', f"{results['recall']:.4f}"],
            ['F1-Score', f"{results['f1_score']:.4f}"],
            ['AUC', f"{results['roc_auc']:.4f}"]
        ]
        
        row = 4
        for metric_name, metric_value in metrics_data:
            ws[f'A{row}'] = metric_name
            ws[f'B{row}'] = metric_value
            ws[f'A{row}'].alignment = cell_alignment
            ws[f'B{row}'].alignment = cell_alignment
            ws[f'A{row}'].border = border
            ws[f'B{row}'].border = border
            row += 1
        
        # Metrics per class
        row += 2
        ws.merge_cells(f'A{row}:D{row}')
        ws[f'A{row}'] = 'METRICS PER CLASS'
        ws[f'A{row}'].font = Font(bold=True, size=12)
        ws[f'A{row}'].alignment = cell_alignment
        ws[f'A{row}'].fill = PatternFill(start_color="663399", end_color="663399", fill_type="solid")
        ws[f'A{row}'].font = Font(bold=True, color="FFFFFF", size=12)
        
        row += 1
        headers = ['Class', 'Precision', 'Recall', 'F1-Score']
        for col, header in enumerate(headers, start=1):
            cell = ws.cell(row=row, column=col)
            cell.value = header
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = cell_alignment
            cell.border = border
        
        row += 1
        class_names = ['Class 0 (Không bệnh)', 'Class 1 (Có bệnh)']
        for i, class_name in enumerate(class_names):
            ws.cell(row=row, column=1, value=class_name)
            ws.cell(row=row, column=2, value=f"{results['precision_per_class'][i]:.4f}")
            ws.cell(row=row, column=3, value=f"{results['recall_per_class'][i]:.4f}")
            ws.cell(row=row, column=4, value=f"{results['f1_per_class'][i]:.4f}")
            
            for col in range(1, 5):
                ws.cell(row=row, column=col).alignment = cell_alignment
                ws.cell(row=row, column=col).border = border
            row += 1
        
        # Model Comparison
        row += 2
        ws.merge_cells(f'A{row}:B{row}')
        ws[f'A{row}'] = 'MODEL ACCURACY COMPARISON'
        ws[f'A{row}'].font = Font(bold=True, size=12)
        ws[f'A{row}'].alignment = cell_alignment
        ws[f'A{row}'].fill = PatternFill(start_color="663399", end_color="663399", fill_type="solid")
        ws[f'A{row}'].font = Font(bold=True, color="FFFFFF", size=12)
        
        row += 1
        ws[f'A{row}'] = 'Model'
        ws[f'B{row}'] = 'Accuracy'
        ws[f'A{row}'].fill = header_fill
        ws[f'B{row}'].fill = header_fill
        ws[f'A{row}'].font = header_font
        ws[f'B{row}'].font = header_font
        ws[f'A{row}'].alignment = cell_alignment
        ws[f'B{row}'].alignment = cell_alignment
        
        row += 1
        for model_name, acc in results['individual_accuracies'].items():
            ws[f'A{row}'] = model_name
            ws[f'B{row}'] = f'{acc:.4f}'
            ws[f'A{row}'].alignment = cell_alignment
            ws[f'B{row}'].alignment = cell_alignment
            ws[f'A{row}'].border = border
            ws[f'B{row}'].border = border
            row += 1
        
        ws[f'A{row}'] = 'Ensemble (Soft Voting)'
        ws[f'B{row}'] = f'{results["test_accuracy"]:.4f}'
        ws[f'A{row}'].font = Font(bold=True)
        ws[f'B{row}'].font = Font(bold=True)
        ws[f'A{row}'].alignment = cell_alignment
        ws[f'B{row}'].alignment = cell_alignment
        ws[f'A{row}'].border = border
        ws[f'B{row}'].border = border
        
        # Confusion Matrix
        row += 3
        ws.merge_cells(f'A{row}:C{row}')
        ws[f'A{row}'] = 'CONFUSION MATRIX'
        ws[f'A{row}'].font = Font(bold=True, size=12)
        ws[f'A{row}'].alignment = cell_alignment
        ws[f'A{row}'].fill = PatternFill(start_color="663399", end_color="663399", fill_type="solid")
        ws[f'A{row}'].font = Font(bold=True, color="FFFFFF", size=12)
        
        row += 1
        ws[f'A{row}'] = ''
        ws[f'B{row}'] = 'Predicted 0'
        ws[f'C{row}'] = 'Predicted 1'
        for col in ['A', 'B', 'C']:
            ws[f'{col}{row}'].fill = header_fill
            ws[f'{col}{row}'].font = header_font
            ws[f'{col}{row}'].alignment = cell_alignment
            ws[f'{col}{row}'].border = border
        
        cm = results['confusion_matrix']
        row += 1
        ws[f'A{row}'] = 'Actual 0'
        ws[f'B{row}'] = cm[0][0]
        ws[f'C{row}'] = cm[0][1]
        for col in ['A', 'B', 'C']:
            ws[f'{col}{row}'].alignment = cell_alignment
            ws[f'{col}{row}'].border = border
        ws[f'A{row}'].font = Font(bold=True)
        
        row += 1
        ws[f'A{row}'] = 'Actual 1'
        ws[f'B{row}'] = cm[1][0]
        ws[f'C{row}'] = cm[1][1]
        for col in ['A', 'B', 'C']:
            ws[f'{col}{row}'].alignment = cell_alignment
            ws[f'{col}{row}'].border = border
        ws[f'A{row}'].font = Font(bold=True)
        
        # Điều chỉnh độ rộng cột
        ws.column_dimensions['A'].width = 30
        ws.column_dimensions['B'].width = 15
        ws.column_dimensions['C'].width = 15
        ws.column_dimensions['D'].width = 15
        
        # Lưu file
        wb.save(filename)
        print(f"✓ Đã xuất metrics ra file: {filename}")
        
    def export_metrics_comparison_to_excel(self, results_by_metric, best_metric, filename='ensemble_metrics_comparison.xlsx'):
        """Xuất kết quả so sánh các metrics ra Excel"""
        print(f"\nĐang xuất kết quả so sánh metrics ra file Excel...")

        wb = Workbook()
        ws = wb.active
        ws.title = "Metrics Comparison"

        # Định dạng style
        header_fill = PatternFill(start_color="70AD47", end_color="70AD47", fill_type="solid")
        header_font = Font(bold=True, color="FFFFFF", size=12)
        best_fill = PatternFill(start_color="FFFF00", end_color="FFFF00", fill_type="solid")  # Yellow for best
        cell_alignment = Alignment(horizontal="center", vertical="center")
        border = Border(
            left=Side(style='thin'),
            right=Side(style='thin'),
            top=Side(style='thin'),
            bottom=Side(style='thin')
        )

        # Tiêu đề chính
        ws.merge_cells('A1:D1')
        ws['A1'] = 'ENSEMBLE METRICS COMPARISON'
        ws['A1'].font = Font(bold=True, size=14)
        ws['A1'].alignment = cell_alignment
        ws['A1'].fill = PatternFill(start_color="385723", end_color="385723", fill_type="solid")
        ws['A1'].font = Font(bold=True, color="FFFFFF", size=14)

        # Header
        headers = ['Metric', 'Ensemble Accuracy', 'Weights', 'Best Model']
        for col, header in enumerate(headers, start=1):
            cell = ws.cell(row=3, column=col)
            cell.value = header
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = cell_alignment
            cell.border = border

        # Dữ liệu
        row = 4
        for metric, data in results_by_metric.items():
            # Metric name
            ws.cell(row=row, column=1, value=metric.upper())
            ws.cell(row=row, column=1).alignment = cell_alignment
            ws.cell(row=row, column=1).border = border

            # Ensemble accuracy
            acc_cell = ws.cell(row=row, column=2, value=f"{data['accuracy']:.4f}")
            acc_cell.alignment = cell_alignment
            acc_cell.border = border
            if metric == best_metric:
                acc_cell.fill = best_fill
                acc_cell.font = Font(bold=True)

            # Weights
            weights_str = ", ".join([f"{name}: {w:.3f}" for name, w in zip(self.model_names, data['weights'])])
            ws.cell(row=row, column=3, value=weights_str)
            ws.cell(row=row, column=3).alignment = Alignment(horizontal="left", vertical="center")
            ws.cell(row=row, column=3).border = border

            # Best individual model
            best_individual = max(data['results']['individual_accuracies'].values())
            ws.cell(row=row, column=4, value=f"{best_individual:.4f}")
            ws.cell(row=row, column=4).alignment = cell_alignment
            ws.cell(row=row, column=4).border = border

            row += 1

        # Summary
        row += 2
        ws.merge_cells(f'A{row}:D{row}')
        ws[f'A{row}'] = f'BEST METRIC: {best_metric.upper()} (Accuracy: {results_by_metric[best_metric]["accuracy"]:.4f})'
        ws[f'A{row}'].font = Font(bold=True, size=12, color="FF0000")
        ws[f'A{row}'].alignment = cell_alignment
        ws[f'A{row}'].fill = PatternFill(start_color="FFFF00", end_color="FFFF00", fill_type="solid")

        # Điều chỉnh độ rộng cột
        ws.column_dimensions['A'].width = 15
        ws.column_dimensions['B'].width = 18
        ws.column_dimensions['C'].width = 50
        ws.column_dimensions['D'].width = 15

        # Lưu file
        wb.save(filename)
        print(f"✓ Đã xuất kết quả so sánh metrics ra file: {filename}")

    def plot_metrics_comparison(self, results_by_metric, best_metric, save_path='ensemble_metrics_comparison.png'):
        """Vẽ biểu đồ so sánh các metrics"""
        metrics = list(results_by_metric.keys())
        accuracies = [results_by_metric[m]['accuracy'] for m in metrics]

        plt.figure(figsize=(12, 6))

        # Bar chart
        bars = plt.bar(metrics, accuracies, color=['lightblue']*len(metrics), alpha=0.7)

        # Highlight best metric
        best_idx = metrics.index(best_metric)
        bars[best_idx].set_color('orange')
        bars[best_idx].set_alpha(1.0)

        # Add value labels on bars
        for bar, acc in zip(bars, accuracies):
            plt.text(bar.get_x() + bar.get_width()/2, bar.get_height() + 0.001,
                    f'{acc:.4f}', ha='center', va='bottom', fontweight='bold')

        plt.xlabel('Weighting Metric', fontsize=12)
        plt.ylabel('Ensemble Accuracy', fontsize=12)
        plt.title('Ensemble Accuracy by Different Weighting Strategies', fontsize=14, fontweight='bold')
        plt.ylim(min(accuracies) - 0.01, max(accuracies) + 0.02)
        plt.grid(axis='y', alpha=0.3)

        # Add legend
        plt.legend([bars[best_idx]], [f'Best: {best_metric.upper()}'], loc='upper right')

        plt.tight_layout()
        plt.savefig(save_path, dpi=300, bbox_inches='tight')
        print(f"Đã lưu biểu đồ so sánh metrics tại: {save_path}")
        plt.close()
    
    def get_ensemble_metrics(self, results):
        """
        Trả về dictionary chứa tất cả metrics chính của ensemble

        Parameters:
        -----------
        results : dict
            Kết quả từ hàm evaluate()

        Returns:
        --------
        dict : Dictionary chứa các metrics chính
        """
        return {
            'accuracy': results['test_accuracy'],
            'precision': results['precision'],
            'recall': results['recall'],
            'f1_score': results['f1_score'],
            'roc_auc': results['roc_auc'],
            'precision_per_class': results['precision_per_class'],
            'recall_per_class': results['recall_per_class'],
            'f1_per_class': results['f1_per_class']
        }

    def print_ensemble_metrics(self, results, title="ENSEMBLE METRICS"):
        """
        In ra các metrics chính của ensemble một cách rõ ràng

        Parameters:
        -----------
        results : dict
            Kết quả từ hàm evaluate()
        title : str
            Tiêu đề hiển thị
        """
        print(f"\n{'='*60}")
        print(f"{title}")
        print(f"{'='*60}")

        # Main metrics
        print(f"Accuracy:  {results['test_accuracy']:.4f} ({results['test_accuracy']*100:.2f}%)")
        print(f"Precision: {results['precision']:.4f}")
        print(f"Recall:    {results['recall']:.4f}")
        print(f"F1-Score:  {results['f1_score']:.4f}")
        print(f"ROC-AUC:   {results['roc_auc']:.4f}")

        # Per-class metrics
        print(f"\nPer-Class Metrics:")
        print(f"Class 0 (Không bệnh):")
        print(f"  Precision: {results['precision_per_class'][0]:.4f}")
        print(f"  Recall:    {results['recall_per_class'][0]:.4f}")
        print(f"  F1-Score:  {results['f1_per_class'][0]:.4f}")

        print(f"Class 1 (Có bệnh):")
        print(f"  Precision: {results['precision_per_class'][1]:.4f}")
        print(f"  Recall:    {results['recall_per_class'][1]:.4f}")
        print(f"  F1-Score:  {results['f1_per_class'][1]:.4f}")

        print(f"{'='*60}")

    def auto_weight_by_metrics(self, metric='accuracy', average='weighted'):
        """
        Tự động gán trọng số cho từng model dựa trên metrics khác nhau

        Parameters:
        -----------
        metric : str, default='accuracy'
            Metric để tính trọng số: 'accuracy', 'f1', 'precision', 'recall', 'roc_auc'
        average : str, default='weighted'
            Cách tính cho multi-class: 'weighted', 'macro', 'micro'
        """
        if self.X_test is None or self.y_test is None:
            print("Chưa có dữ liệu test để tính trọng số!")
            return

        print(f"\nĐang tính trọng số cho từng model dựa trên {metric}...")

        scores = []
        for i, (model, scaler) in enumerate(zip(self.models, self.scalers)):
            if scaler is not None:
                X_scaled = scaler.transform(self.X_test)
            else:
                X_scaled = self.X_test

            pred = model.predict(X_scaled)

            if metric == 'accuracy':
                score = accuracy_score(self.y_test, pred)
            elif metric == 'f1':
                score = f1_score(self.y_test, pred, average=average)
            elif metric == 'precision':
                score = precision_score(self.y_test, pred, average=average)
            elif metric == 'recall':
                score = recall_score(self.y_test, pred, average=average)
            elif metric == 'roc_auc':
                try:
                    proba = model.predict_proba(X_scaled)[:, 1]
                    score = roc_auc_score(self.y_test, proba)
                except:
                    print(f"Model {self.model_names[i]} không hỗ trợ predict_proba, dùng accuracy thay thế")
                    score = accuracy_score(self.y_test, pred)
            else:
                raise ValueError(f"Metric '{metric}' không được hỗ trợ")

            scores.append(score)
            print(f"Model {self.model_names[i]}: {metric} = {score:.4f}")

        # Normalize scores thành weights
        total = sum(scores)
        if total == 0:
            self.weights = [1.0/len(scores)] * len(scores)
        else:
            self.weights = [score/total for score in scores]

        print(f"Trọng số mới dựa trên {metric}: {dict(zip(self.model_names, self.weights))}")

    def auto_weight_by_accuracy(self):
        """Alias cho backward compatibility"""
        self.auto_weight_by_metrics(metric='accuracy')
