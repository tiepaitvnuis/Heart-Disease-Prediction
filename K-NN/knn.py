import pandas as pd
import numpy as np
from sklearn.model_selection import train_test_split, cross_val_score, GridSearchCV
from sklearn.preprocessing import StandardScaler
from sklearn.neighbors import KNeighborsClassifier
from sklearn.metrics import accuracy_score, classification_report, confusion_matrix, roc_auc_score, roc_curve, precision_score, recall_score, f1_score
import matplotlib.pyplot as plt
import seaborn as sns
import joblib
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
import warnings
import os
warnings.filterwarnings('ignore')


class HeartDiseaseKNN:
    def __init__(self, n_neighbors=5, weights='uniform', metric='minkowski'):
        """
        Khởi tạo model KNN cho dự đoán bệnh tim
        
        Parameters:
        -----------
        n_neighbors : int, default=5
            Số lượng láng giềng gần nhất
        weights : {'uniform', 'distance'}, default='uniform'
            Trọng số cho việc dự đoán
        metric : str, default='minkowski'
            Độ đo khoảng cách
        """
        self.n_neighbors = n_neighbors
        self.weights = weights
        self.metric = metric
        self.model = KNeighborsClassifier(n_neighbors=n_neighbors, 
                                         weights=weights, 
                                         metric=metric)
        self.scaler = StandardScaler()
        self.X_train = None
        self.X_test = None
        self.y_train = None
        self.y_test = None
        
    def load_data(self, file_path):
        """Tải dữ liệu từ file CSV"""
        print(f"Đang tải dữ liệu từ {file_path}...")
        self.data = pd.read_csv(file_path)
        print(f"Dữ liệu đã tải: {self.data.shape[0]} hàng, {self.data.shape[1]} cột")
        return self.data
    
    def preprocess_data(self, test_size=0.2, random_state=42):
        """Tiền xử lý dữ liệu"""
        print("\nĐang tiền xử lý dữ liệu...")
        
        # Tách features và target
        X = self.data.drop('target', axis=1)
        y = self.data['target']
        
        # Chia dữ liệu train/test
        self.X_train, self.X_test, self.y_train, self.y_test = train_test_split(
            X, y, test_size=test_size, random_state=random_state, stratify=y
        )
        
        # Chuẩn hóa dữ liệu (quan trọng cho KNN)
        self.X_train = self.scaler.fit_transform(self.X_train)
        self.X_test = self.scaler.transform(self.X_test)
        
        print(f"Dữ liệu train: {self.X_train.shape}")
        print(f"Dữ liệu test: {self.X_test.shape}")
        
    def train(self):
        """Huấn luyện model KNN"""
        print("\nĐang huấn luyện model KNN...")
        self.model.fit(self.X_train, self.y_train)
        print("Huấn luyện hoàn tất!")
        
    def evaluate(self):
        """Đánh giá model"""
        print("\n" + "="*60)
        print("KẾT QUẢ ĐÁNH GIÁ MODEL KNN")
        print("="*60)
        
        # Dự đoán
        y_train_pred = self.model.predict(self.X_train)
        y_test_pred = self.model.predict(self.X_test)
        
        # Tính toán các metrics
        train_accuracy = accuracy_score(self.y_train, y_train_pred)
        test_accuracy = accuracy_score(self.y_test, y_test_pred)
        
        # Precision, Recall, F1-Score cho test set
        precision = precision_score(self.y_test, y_test_pred, average='weighted')
        recall = recall_score(self.y_test, y_test_pred, average='weighted')
        f1 = f1_score(self.y_test, y_test_pred, average='weighted')
        
        # Precision, Recall, F1 cho từng class
        precision_per_class = precision_score(self.y_test, y_test_pred, average=None)
        recall_per_class = recall_score(self.y_test, y_test_pred, average=None)
        f1_per_class = f1_score(self.y_test, y_test_pred, average=None)
        
        print(f"\nĐộ chính xác trên tập train: {train_accuracy:.4f} ({train_accuracy*100:.2f}%)")
        print(f"Độ chính xác trên tập test: {test_accuracy:.4f} ({test_accuracy*100:.2f}%)")
        print(f"Precision (weighted): {precision:.4f}")
        print(f"Recall (weighted): {recall:.4f}")
        print(f"F1-Score (weighted): {f1:.4f}")
        
        # Cross-validation
        cv_scores = cross_val_score(self.model, self.X_train, self.y_train, cv=5)
        print(f"\nCross-Validation Scores: {cv_scores}")
        print(f"CV Mean: {cv_scores.mean():.4f} (+/- {cv_scores.std() * 2:.4f})")
        
        # Classification report
        print("\n" + "-"*60)
        print("CLASSIFICATION REPORT (Test Set):")
        print("-"*60)
        print(classification_report(self.y_test, y_test_pred, 
                                   target_names=['Không bệnh (0)', 'Có bệnh (1)']))
        
        # Confusion Matrix
        cm = confusion_matrix(self.y_test, y_test_pred)
        print("\nConfusion Matrix:")
        print(cm)
        
        # ROC AUC Score
        y_test_proba = self.model.predict_proba(self.X_test)[:, 1]
        roc_auc = roc_auc_score(self.y_test, y_test_proba)
        print(f"\nROC AUC Score: {roc_auc:.4f}")
        
        return {
            'train_accuracy': train_accuracy,
            'test_accuracy': test_accuracy,
            'precision': precision,
            'recall': recall,
            'f1_score': f1,
            'precision_per_class': precision_per_class,
            'recall_per_class': recall_per_class,
            'f1_per_class': f1_per_class,
            'cv_scores': cv_scores,
            'roc_auc': roc_auc,
            'confusion_matrix': cm
        }
    
    def plot_confusion_matrix(self, save_path='knn_confusion_matrix.png'):
        """Vẽ confusion matrix"""
        y_pred = self.model.predict(self.X_test)
        cm = confusion_matrix(self.y_test, y_pred)
        
        plt.figure(figsize=(8, 6))
        sns.heatmap(cm, annot=True, fmt='d', cmap='Blues', 
                   xticklabels=['Không bệnh', 'Có bệnh'],
                   yticklabels=['Không bệnh', 'Có bệnh'])
        plt.title('Confusion Matrix - KNN Model')
        plt.ylabel('Actual')
        plt.xlabel('Predicted')
        plt.tight_layout()
        plt.savefig(save_path, dpi=300, bbox_inches='tight')
        print(f"\nĐã lưu confusion matrix tại: {save_path}")
        plt.close()
        
    def plot_roc_curve(self, save_path='knn_roc_curve.png'):
        """Vẽ ROC curve"""
        y_proba = self.model.predict_proba(self.X_test)[:, 1]
        fpr, tpr, thresholds = roc_curve(self.y_test, y_proba)
        roc_auc = roc_auc_score(self.y_test, y_proba)
        
        plt.figure(figsize=(8, 6))
        plt.plot(fpr, tpr, color='darkorange', lw=2, 
                label=f'ROC curve (AUC = {roc_auc:.4f})')
        plt.plot([0, 1], [0, 1], color='navy', lw=2, linestyle='--', label='Random')
        plt.xlim([0.0, 1.0])
        plt.ylim([0.0, 1.05])
        plt.xlabel('False Positive Rate')
        plt.ylabel('True Positive Rate')
        plt.title('ROC Curve - KNN Model')
        plt.legend(loc="lower right")
        plt.grid(alpha=0.3)
        plt.tight_layout()
        plt.savefig(save_path, dpi=300, bbox_inches='tight')
        print(f"Đã lưu ROC curve tại: {save_path}")
        plt.close()
        
    def find_best_k(self, k_range=range(1, 31)):
        """Tìm giá trị K tốt nhất"""
        print("\nĐang tìm giá trị K tốt nhất...")
        
        train_scores = []
        test_scores = []
        
        for k in k_range:
            knn = KNeighborsClassifier(n_neighbors=k)
            knn.fit(self.X_train, self.y_train)
            
            train_scores.append(knn.score(self.X_train, self.y_train))
            test_scores.append(knn.score(self.X_test, self.y_test))
        
        # Vẽ biểu đồ
        plt.figure(figsize=(10, 6))
        plt.plot(k_range, train_scores, marker='o', label='Train Score', linewidth=2)
        plt.plot(k_range, test_scores, marker='s', label='Test Score', linewidth=2)
        plt.xlabel('Number of Neighbors (K)')
        plt.ylabel('Accuracy')
        plt.title('KNN: Accuracy vs K Value')
        plt.legend()
        plt.grid(alpha=0.3)
        plt.tight_layout()
        plt.savefig('knn_k_tuning.png', dpi=300, bbox_inches='tight')
        print("Đã lưu biểu đồ K tuning tại: knn_k_tuning.png")
        plt.close()
        
        # Tìm K tốt nhất
        best_k = k_range[test_scores.index(max(test_scores))]
        print(f"\nGiá trị K tốt nhất: {best_k}")
        print(f"Độ chính xác tốt nhất: {max(test_scores):.4f}")
        
        return best_k
    
    def grid_search(self):
        """Tìm kiếm tham số tốt nhất bằng GridSearch"""
        print("\nĐang thực hiện GridSearch...")
        
        param_grid = {
            'n_neighbors': [3, 5, 7, 9, 11, 13, 15],
            'weights': ['uniform', 'distance'],
            'metric': ['euclidean', 'manhattan', 'minkowski']
        }
        
        grid_search = GridSearchCV(
            KNeighborsClassifier(),
            param_grid,
            cv=5,
            scoring='accuracy',
            n_jobs=-1,
            verbose=1
        )
        
        grid_search.fit(self.X_train, self.y_train)
        
        print("\nTham số tốt nhất:", grid_search.best_params_)
        print(f"Độ chính xác tốt nhất: {grid_search.best_score_:.4f}")
        
        # Cập nhật model với tham số tốt nhất
        self.model = grid_search.best_estimator_
        
        return grid_search.best_params_
    
    def predict(self, X):
        """Dự đoán cho dữ liệu mới"""
        X_scaled = self.scaler.transform(X)
        predictions = self.model.predict(X_scaled)
        probabilities = self.model.predict_proba(X_scaled)
        return predictions, probabilities
    
    def save_model(self, model_path='knn_heart_disease_model.pkl', 
                   scaler_path='knn_scaler.pkl'):
        """Lưu model và scaler"""
        joblib.dump(self.model, model_path)
        joblib.dump(self.scaler, scaler_path)
        print(f"\nĐã lưu model tại: {model_path}")
        print(f"Đã lưu scaler tại: {scaler_path}")
        
    def load_model(self, model_path='knn_heart_disease_model.pkl', 
                   scaler_path='knn_scaler.pkl'):
        """Tải model và scaler"""
        self.model = joblib.load(model_path)
        self.scaler = joblib.load(scaler_path)
        print(f"Đã tải model từ: {model_path}")
        print(f"Đã tải scaler từ: {scaler_path}")
    
    def export_metrics_to_excel(self, results, filename='knn_metrics.xlsx'):
        """Xuất các metrics ra file Excel"""
        print(f"\nĐang xuất kết quả ra file Excel...")
        
        wb = Workbook()
        ws = wb.active
        ws.title = "KNN Metrics"
        
        # Định dạng style
        header_fill = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
        header_font = Font(bold=True, color="FFFFFF", size=12)
        cell_alignment = Alignment(horizontal="center", vertical="center")
        border = Border(
            left=Side(style='thin'),
            right=Side(style='thin'),
            top=Side(style='thin'),
            bottom=Side(style='thin')
        )
        
        # Tiêu đề chính
        ws.merge_cells('A1:B1')
        ws['A1'] = 'KNN MODEL - EVALUATION METRICS'
        ws['A1'].font = Font(bold=True, size=14)
        ws['A1'].alignment = cell_alignment
        ws['A1'].fill = PatternFill(start_color="203864", end_color="203864", fill_type="solid")
        ws['A1'].font = Font(bold=True, color="FFFFFF", size=14)
        
        # Header cho metrics tổng quan
        ws['A3'] = 'Metric'
        ws['B3'] = 'Value'
        ws['A3'].fill = header_fill
        ws['B3'].fill = header_fill
        ws['A3'].font = header_font
        ws['B3'].font = header_font
        ws['A3'].alignment = cell_alignment
        ws['B3'].alignment = cell_alignment
        
        # Dữ liệu metrics
        metrics_data = [
            ['Accuracy', f"{results['test_accuracy']:.4f}"],
            ['Precision', f"{results['precision']:.4f}"],
            ['Recall', f"{results['recall']:.4f}"],
            ['F1-Score', f"{results['f1_score']:.4f}"],
            ['AUC', f"{results['roc_auc']:.4f}"]
        ]
        
        row = 4
        for metric_name, metric_value in metrics_data:
            ws[f'A{row}'] = metric_name
            ws[f'B{row}'] = metric_value
            ws[f'A{row}'].alignment = cell_alignment
            ws[f'B{row}'].alignment = cell_alignment
            ws[f'A{row}'].border = border
            ws[f'B{row}'].border = border
            row += 1
        
        # Metrics per class
        row += 2
        ws.merge_cells(f'A{row}:D{row}')
        ws[f'A{row}'] = 'METRICS PER CLASS'
        ws[f'A{row}'].font = Font(bold=True, size=12)
        ws[f'A{row}'].alignment = cell_alignment
        ws[f'A{row}'].fill = PatternFill(start_color="203864", end_color="203864", fill_type="solid")
        ws[f'A{row}'].font = Font(bold=True, color="FFFFFF", size=12)
        
        row += 1
        headers = ['Class', 'Precision', 'Recall', 'F1-Score']
        for col, header in enumerate(headers, start=1):
            cell = ws.cell(row=row, column=col)
            cell.value = header
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = cell_alignment
            cell.border = border
        
        row += 1
        class_names = ['Class 0 (Không bệnh)', 'Class 1 (Có bệnh)']
        for i, class_name in enumerate(class_names):
            ws.cell(row=row, column=1, value=class_name)
            ws.cell(row=row, column=2, value=f"{results['precision_per_class'][i]:.4f}")
            ws.cell(row=row, column=3, value=f"{results['recall_per_class'][i]:.4f}")
            ws.cell(row=row, column=4, value=f"{results['f1_per_class'][i]:.4f}")
            
            for col in range(1, 5):
                ws.cell(row=row, column=col).alignment = cell_alignment
                ws.cell(row=row, column=col).border = border
            row += 1
        
        # Cross-Validation Scores
        row += 2
        ws.merge_cells(f'A{row}:B{row}')
        ws[f'A{row}'] = 'CROSS-VALIDATION SCORES'
        ws[f'A{row}'].font = Font(bold=True, size=12)
        ws[f'A{row}'].alignment = cell_alignment
        ws[f'A{row}'].fill = PatternFill(start_color="203864", end_color="203864", fill_type="solid")
        ws[f'A{row}'].font = Font(bold=True, color="FFFFFF", size=12)
        
        row += 1
        ws[f'A{row}'] = 'Fold'
        ws[f'B{row}'] = 'Score'
        ws[f'A{row}'].fill = header_fill
        ws[f'B{row}'].fill = header_fill
        ws[f'A{row}'].font = header_font
        ws[f'B{row}'].font = header_font
        ws[f'A{row}'].alignment = cell_alignment
        ws[f'B{row}'].alignment = cell_alignment
        
        row += 1
        for i, score in enumerate(results['cv_scores'], 1):
            ws[f'A{row}'] = f'Fold {i}'
            ws[f'B{row}'] = f'{score:.4f}'
            ws[f'A{row}'].alignment = cell_alignment
            ws[f'B{row}'].alignment = cell_alignment
            ws[f'A{row}'].border = border
            ws[f'B{row}'].border = border
            row += 1
        
        ws[f'A{row}'] = 'Mean'
        ws[f'B{row}'] = f'{results["cv_scores"].mean():.4f}'
        ws[f'A{row}'].font = Font(bold=True)
        ws[f'B{row}'].font = Font(bold=True)
        ws[f'A{row}'].alignment = cell_alignment
        ws[f'B{row}'].alignment = cell_alignment
        ws[f'A{row}'].border = border
        ws[f'B{row}'].border = border
        
        row += 1
        ws[f'A{row}'] = 'Std Dev'
        ws[f'B{row}'] = f'{results["cv_scores"].std():.4f}'
        ws[f'A{row}'].font = Font(bold=True)
        ws[f'B{row}'].font = Font(bold=True)
        ws[f'A{row}'].alignment = cell_alignment
        ws[f'B{row}'].alignment = cell_alignment
        ws[f'A{row}'].border = border
        ws[f'B{row}'].border = border
        
        # Confusion Matrix
        row += 3
        ws.merge_cells(f'A{row}:C{row}')
        ws[f'A{row}'] = 'CONFUSION MATRIX'
        ws[f'A{row}'].font = Font(bold=True, size=12)
        ws[f'A{row}'].alignment = cell_alignment
        ws[f'A{row}'].fill = PatternFill(start_color="203864", end_color="203864", fill_type="solid")
        ws[f'A{row}'].font = Font(bold=True, color="FFFFFF", size=12)
        
        row += 1
        ws[f'A{row}'] = ''
        ws[f'B{row}'] = 'Predicted 0'
        ws[f'C{row}'] = 'Predicted 1'
        for col in ['A', 'B', 'C']:
            ws[f'{col}{row}'].fill = header_fill
            ws[f'{col}{row}'].font = header_font
            ws[f'{col}{row}'].alignment = cell_alignment
            ws[f'{col}{row}'].border = border
        
        cm = results['confusion_matrix']
        row += 1
        ws[f'A{row}'] = 'Actual 0'
        ws[f'B{row}'] = cm[0][0]
        ws[f'C{row}'] = cm[0][1]
        for col in ['A', 'B', 'C']:
            ws[f'{col}{row}'].alignment = cell_alignment
            ws[f'{col}{row}'].border = border
        ws[f'A{row}'].font = Font(bold=True)
        
        row += 1
        ws[f'A{row}'] = 'Actual 1'
        ws[f'B{row}'] = cm[1][0]
        ws[f'C{row}'] = cm[1][1]
        for col in ['A', 'B', 'C']:
            ws[f'{col}{row}'].alignment = cell_alignment
            ws[f'{col}{row}'].border = border
        ws[f'A{row}'].font = Font(bold=True)
        
        # Điều chỉnh độ rộng cột
        ws.column_dimensions['A'].width = 25
        ws.column_dimensions['B'].width = 15
        ws.column_dimensions['C'].width = 15
        ws.column_dimensions['D'].width = 15
        
        # Lưu file
        wb.save(filename)
        print(f"✓ Đã xuất metrics ra file: {filename}")
        
        return filename


def main():
    """Hàm chính để huấn luyện và đánh giá model KNN"""
    
    # Tạo thư mục output
    output_dir = 'knn_results'
    if not os.path.exists(output_dir):
        os.makedirs(output_dir)
    
    print("="*60)
    print("HUẤN LUYỆN MODEL KNN CHO DỰ ĐOÁN BỆNH TIM")
    print("="*60)
    
    # Khởi tạo model
    knn = HeartDiseaseKNN(n_neighbors=5, weights='uniform', metric='minkowski')
    
    # Tải dữ liệu
    data_file = 'Heart_disease_cleveland_new(2).csv'
    knn.load_data(data_file)
    
    # Tiền xử lý (85% train, 15% test)
    knn.preprocess_data(test_size=0.15, random_state=42)
    
    # Tìm K tốt nhất
    best_k = knn.find_best_k(k_range=range(1, 31))
    
    # Cập nhật model với K tốt nhất
    print(f"\nCập nhật model với K = {best_k}")
    knn.n_neighbors = best_k
    knn.model = KNeighborsClassifier(n_neighbors=best_k, 
                                     weights=knn.weights, 
                                     metric=knn.metric)
    
    # Huấn luyện model
    knn.train()
    
    # Đánh giá model
    results = knn.evaluate()
    
    # Vẽ các biểu đồ
    knn.plot_confusion_matrix(save_path=os.path.join(output_dir, 'knn_confusion_matrix.png'))
    knn.plot_roc_curve(save_path=os.path.join(output_dir, 'knn_roc_curve.png'))
    
    # Tùy chọn: Thực hiện GridSearch để tìm tham số tốt nhất
    use_gridsearch = True
    if use_gridsearch:
        best_params = knn.grid_search()
        
        # Đánh giá lại với tham số tốt nhất
        print("\nĐánh giá lại model với tham số tốt nhất...")
        results = knn.evaluate()
        knn.plot_confusion_matrix(save_path=os.path.join(output_dir, 'knn_confusion_matrix_optimized.png'))
        knn.plot_roc_curve(save_path=os.path.join(output_dir, 'knn_roc_curve_optimized.png'))
    
    # Lưu model sau khi train
    knn.save_model(
        model_path=os.path.join(output_dir, 'knn_heart_disease_model.pkl'),
        scaler_path=os.path.join(output_dir, 'knn_scaler.pkl')
    )
    
    # Xuất metrics ra Excel
    knn.export_metrics_to_excel(results, filename=os.path.join(output_dir, 'knn_metrics.xlsx'))
    
    print("\n" + "="*60)
    print("HOÀN THÀNH!")
    print("="*60)
    print("\nKết quả:")
    print(f"- Độ chính xác trên test set: {results['test_accuracy']*100:.2f}%")
    print(f"- Precision: {results['precision']:.4f}")
    print(f"- Recall: {results['recall']:.4f}")
    print(f"- F1-Score: {results['f1_score']:.4f}")
    print(f"- ROC AUC Score: {results['roc_auc']:.4f}")
    print(f"- Cross-Validation Mean: {results['cv_scores'].mean():.4f}")
    
    # Demo dự đoán
    print("\n" + "="*60)
    print("DEMO DỰ ĐOÁN")
    print("="*60)
    
    # Lấy một mẫu từ test set
    sample = pd.DataFrame([knn.X_test[0]])
    prediction, probability = knn.predict(sample)
    
    print(f"\nMẫu dự đoán:")
    print(f"Kết quả dự đoán: {'Có bệnh tim' if prediction[0] == 1 else 'Không có bệnh tim'}")
    print(f"Xác suất không bệnh: {probability[0][0]:.4f}")
    print(f"Xác suất có bệnh: {probability[0][1]:.4f}")


if __name__ == "__main__":
    main()
