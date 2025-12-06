import pandas as pd
import numpy as np
from sklearn.model_selection import train_test_split, cross_val_score, GridSearchCV
from sklearn.preprocessing import StandardScaler
from sklearn.tree import DecisionTreeClassifier, plot_tree
from sklearn.metrics import accuracy_score, classification_report, confusion_matrix, roc_auc_score, roc_curve, precision_score, recall_score, f1_score
import matplotlib.pyplot as plt
import seaborn as sns
import joblib
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
import warnings
warnings.filterwarnings('ignore')


class HeartDiseaseDecisionTree:
    def __init__(self, max_depth=None, min_samples_split=2, min_samples_leaf=1, criterion='gini', random_state=42):
        """
        Khởi tạo model Decision Tree cho dự đoán bệnh tim
        
        Parameters:
        -----------
        max_depth : int, default=None
            Độ sâu tối đa của cây
        min_samples_split : int, default=2
            Số mẫu tối thiểu để chia node
        min_samples_leaf : int, default=1
            Số mẫu tối thiểu ở mỗi lá
        criterion : {'gini', 'entropy'}, default='gini'
            Hàm đo độ tinh khiết
        random_state : int, default=42
            Random seed
        """
        self.max_depth = max_depth
        self.min_samples_split = min_samples_split
        self.min_samples_leaf = min_samples_leaf
        self.criterion = criterion
        self.random_state = random_state
        
        self.model = DecisionTreeClassifier(
            max_depth=max_depth,
            min_samples_split=min_samples_split,
            min_samples_leaf=min_samples_leaf,
            criterion=criterion,
            random_state=random_state
        )
        self.scaler = StandardScaler()
        self.X_train = None
        self.X_test = None
        self.y_train = None
        self.y_test = None
        self.feature_names = None
        
    def load_data(self, file_path):
        """Tải dữ liệu từ file CSV"""
        print(f"Đang tải dữ liệu từ {file_path}...")
        self.data = pd.read_csv(file_path)
        print(f"Dữ liệu đã tải: {self.data.shape[0]} hàng, {self.data.shape[1]} cột")
        return self.data
    
    def preprocess_data(self, test_size=0.15, random_state=42):
        """Tiền xử lý dữ liệu"""
        print("\nĐang tiền xử lý dữ liệu...")
        
        # Tách features và target
        X = self.data.drop('target', axis=1)
        y = self.data['target']
        self.feature_names = X.columns.tolist()
        
        # Chia dữ liệu train/test
        self.X_train, self.X_test, self.y_train, self.y_test = train_test_split(
            X, y, test_size=test_size, random_state=random_state, stratify=y
        )
        
        # Chuẩn hóa dữ liệu
        self.X_train = self.scaler.fit_transform(self.X_train)
        self.X_test = self.scaler.transform(self.X_test)
        
        print(f"Dữ liệu train: {self.X_train.shape}")
        print(f"Dữ liệu test: {self.X_test.shape}")
        
    def train(self):
        """Huấn luyện model Decision Tree"""
        print("\nĐang huấn luyện model Decision Tree...")
        self.model.fit(self.X_train, self.y_train)
        print("Huấn luyện hoàn tất!")
        
    def evaluate(self):
        """Đánh giá model"""
        print("\n" + "="*60)
        print("KẾT QUẢ ĐÁNH GIÁ MODEL DECISION TREE")
        print("="*60)
        
        # Dự đoán
        y_train_pred = self.model.predict(self.X_train)
        y_test_pred = self.model.predict(self.X_test)
        
        # Tính toán các metrics
        train_accuracy = accuracy_score(self.y_train, y_train_pred)
        test_accuracy = accuracy_score(self.y_test, y_test_pred)
        
        # Precision, Recall, F1-Score cho test set
        precision = precision_score(self.y_test, y_test_pred, average='weighted')
        recall = recall_score(self.y_test, y_test_pred, average='weighted')
        f1 = f1_score(self.y_test, y_test_pred, average='weighted')
        
        # Precision, Recall, F1 cho từng class
        precision_per_class = precision_score(self.y_test, y_test_pred, average=None)
        recall_per_class = recall_score(self.y_test, y_test_pred, average=None)
        f1_per_class = f1_score(self.y_test, y_test_pred, average=None)
        
        print(f"\nĐộ chính xác trên tập train: {train_accuracy:.4f} ({train_accuracy*100:.2f}%)")
        print(f"Độ chính xác trên tập test: {test_accuracy:.4f} ({test_accuracy*100:.2f}%)")
        print(f"Precision (weighted): {precision:.4f}")
        print(f"Recall (weighted): {recall:.4f}")
        print(f"F1-Score (weighted): {f1:.4f}")
        
        # Cross-validation
        cv_scores = cross_val_score(self.model, self.X_train, self.y_train, cv=5)
        print(f"\nCross-Validation Scores: {cv_scores}")
        print(f"CV Mean: {cv_scores.mean():.4f} (+/- {cv_scores.std() * 2:.4f})")
        
        # Classification report
        print("\n" + "-"*60)
        print("CLASSIFICATION REPORT (Test Set):")
        print("-"*60)
        print(classification_report(self.y_test, y_test_pred, 
                                   target_names=['Không bệnh (0)', 'Có bệnh (1)']))
        
        # Confusion Matrix
        cm = confusion_matrix(self.y_test, y_test_pred)
        print("\nConfusion Matrix:")
        print(cm)
        
        # ROC AUC Score
        y_test_proba = self.model.predict_proba(self.X_test)[:, 1]
        roc_auc = roc_auc_score(self.y_test, y_test_proba)
        print(f"\nROC AUC Score: {roc_auc:.4f}")
        
        # Feature Importance
        feature_importance = pd.DataFrame({
            'feature': self.feature_names,
            'importance': self.model.feature_importances_
        }).sort_values('importance', ascending=False)
        
        print("\n" + "-"*60)
        print("FEATURE IMPORTANCE (Top 10):")
        print("-"*60)
        print(feature_importance.head(10).to_string(index=False))
        
        return {
            'train_accuracy': train_accuracy,
            'test_accuracy': test_accuracy,
            'precision': precision,
            'recall': recall,
            'f1_score': f1,
            'precision_per_class': precision_per_class,
            'recall_per_class': recall_per_class,
            'f1_per_class': f1_per_class,
            'cv_scores': cv_scores,
            'roc_auc': roc_auc,
            'confusion_matrix': cm,
            'feature_importance': feature_importance
        }
    
    def plot_confusion_matrix(self, save_path='dt_confusion_matrix.png'):
        """Vẽ confusion matrix"""
        y_pred = self.model.predict(self.X_test)
        cm = confusion_matrix(self.y_test, y_pred)
        
        plt.figure(figsize=(8, 6))
        sns.heatmap(cm, annot=True, fmt='d', cmap='Greens', 
                   xticklabels=['Không bệnh', 'Có bệnh'],
                   yticklabels=['Không bệnh', 'Có bệnh'])
        plt.title('Confusion Matrix - Decision Tree Model')
        plt.ylabel('Actual')
        plt.xlabel('Predicted')
        plt.tight_layout()
        plt.savefig(save_path, dpi=300, bbox_inches='tight')
        print(f"\nĐã lưu confusion matrix tại: {save_path}")
        plt.close()
        
    def plot_roc_curve(self, save_path='dt_roc_curve.png'):
        """Vẽ ROC curve"""
        y_proba = self.model.predict_proba(self.X_test)[:, 1]
        fpr, tpr, thresholds = roc_curve(self.y_test, y_proba)
        roc_auc = roc_auc_score(self.y_test, y_proba)
        
        plt.figure(figsize=(8, 6))
        plt.plot(fpr, tpr, color='green', lw=2, 
                label=f'ROC curve (AUC = {roc_auc:.4f})')
        plt.plot([0, 1], [0, 1], color='navy', lw=2, linestyle='--', label='Random')
        plt.xlim([0.0, 1.0])
        plt.ylim([0.0, 1.05])
        plt.xlabel('False Positive Rate')
        plt.ylabel('True Positive Rate')
        plt.title('ROC Curve - Decision Tree Model')
        plt.legend(loc="lower right")
        plt.grid(alpha=0.3)
        plt.tight_layout()
        plt.savefig(save_path, dpi=300, bbox_inches='tight')
        print(f"Đã lưu ROC curve tại: {save_path}")
        plt.close()
    
    def plot_feature_importance(self, save_path='dt_feature_importance.png', top_n=10):
        """Vẽ biểu đồ feature importance"""
        feature_importance = pd.DataFrame({
            'feature': self.feature_names,
            'importance': self.model.feature_importances_
        }).sort_values('importance', ascending=False).head(top_n)
        
        plt.figure(figsize=(10, 6))
        plt.barh(range(len(feature_importance)), feature_importance['importance'], color='green')
        plt.yticks(range(len(feature_importance)), feature_importance['feature'])
        plt.xlabel('Importance')
        plt.title(f'Top {top_n} Feature Importance - Decision Tree')
        plt.gca().invert_yaxis()
        plt.tight_layout()
        plt.savefig(save_path, dpi=300, bbox_inches='tight')
        print(f"Đã lưu feature importance tại: {save_path}")
        plt.close()
    
    def plot_tree_visualization(self, save_path='dt_tree_visualization.png', max_depth_display=3):
        """Vẽ cây quyết định"""
        plt.figure(figsize=(20, 10))
        plot_tree(self.model, 
                 feature_names=self.feature_names,
                 class_names=['Không bệnh', 'Có bệnh'],
                 filled=True,
                 rounded=True,
                 max_depth=max_depth_display,
                 fontsize=10)
        plt.title('Decision Tree Visualization')
        plt.tight_layout()
        plt.savefig(save_path, dpi=300, bbox_inches='tight')
        print(f"Đã lưu tree visualization tại: {save_path}")
        plt.close()
        
    def tune_max_depth(self, depth_range=range(1, 21), save_path='dt_depth_tuning.png'):
        """Tìm max_depth tốt nhất"""
        print("\nĐang tìm max_depth tốt nhất...")
        
        train_scores = []
        test_scores = []
        
        for depth in depth_range:
            dt = DecisionTreeClassifier(max_depth=depth, random_state=self.random_state)
            dt.fit(self.X_train, self.y_train)
            
            train_scores.append(dt.score(self.X_train, self.y_train))
            test_scores.append(dt.score(self.X_test, self.y_test))
        
        # Vẽ biểu đồ
        plt.figure(figsize=(10, 6))
        plt.plot(depth_range, train_scores, marker='o', label='Train Score', linewidth=2)
        plt.plot(depth_range, test_scores, marker='s', label='Test Score', linewidth=2)
        plt.xlabel('Max Depth')
        plt.ylabel('Accuracy')
        plt.title('Decision Tree: Accuracy vs Max Depth')
        plt.legend()
        plt.grid(alpha=0.3)
        plt.tight_layout()
        plt.savefig(save_path, dpi=300, bbox_inches='tight')
        print(f"Đã lưu biểu đồ depth tuning tại: {save_path}")
        plt.close()
        
        # Tìm depth tốt nhất
        best_depth = depth_range[test_scores.index(max(test_scores))]
        print(f"\nMax depth tốt nhất: {best_depth}")
        print(f"Độ chính xác tốt nhất: {max(test_scores):.4f}")
        
        return best_depth
    
    def grid_search(self):
        """Tìm kiếm tham số tốt nhất bằng GridSearch"""
        print("\nĐang thực hiện GridSearch...")
        
        param_grid = {
            'max_depth': [3, 5, 7, 9, 11, 15, 20, None],
            'min_samples_split': [2, 5, 10, 20],
            'min_samples_leaf': [1, 2, 4, 8],
            'criterion': ['gini', 'entropy']
        }
        
        grid_search = GridSearchCV(
            DecisionTreeClassifier(random_state=self.random_state),
            param_grid,
            cv=5,
            scoring='accuracy',
            n_jobs=-1,
            verbose=1
        )
        
        grid_search.fit(self.X_train, self.y_train)
        
        print("\nTham số tốt nhất:", grid_search.best_params_)
        print(f"Độ chính xác tốt nhất: {grid_search.best_score_:.4f}")
        
        # Cập nhật model với tham số tốt nhất
        self.model = grid_search.best_estimator_
        
        return grid_search.best_params_
    
    def predict(self, X):
        """Dự đoán cho dữ liệu mới"""
        X_scaled = self.scaler.transform(X)
        predictions = self.model.predict(X_scaled)
        probabilities = self.model.predict_proba(X_scaled)
        return predictions, probabilities
    
    def save_model(self, model_path='dt_heart_disease_model.pkl', 
                   scaler_path='dt_scaler.pkl'):
        """Lưu model và scaler"""
        joblib.dump(self.model, model_path)
        joblib.dump(self.scaler, scaler_path)
        print(f"\nĐã lưu model tại: {model_path}")
        print(f"Đã lưu scaler tại: {scaler_path}")
        
    def load_model(self, model_path='dt_heart_disease_model.pkl', 
                   scaler_path='dt_scaler.pkl'):
        """Tải model và scaler"""
        self.model = joblib.load(model_path)
        self.scaler = joblib.load(scaler_path)
        print(f"Đã tải model từ: {model_path}")
        print(f"Đã tải scaler từ: {scaler_path}")
    
    def export_metrics_to_excel(self, results, filename='dt_metrics.xlsx'):
        """Xuất các metrics ra file Excel"""
        print(f"\nĐang xuất kết quả ra file Excel...")
        
        wb = Workbook()
        ws = wb.active
        ws.title = "Decision Tree Metrics"
        
        # Định dạng style
        header_fill = PatternFill(start_color="70AD47", end_color="70AD47", fill_type="solid")
        header_font = Font(bold=True, color="FFFFFF", size=12)
        cell_alignment = Alignment(horizontal="center", vertical="center")
        border = Border(
            left=Side(style='thin'),
            right=Side(style='thin'),
            top=Side(style='thin'),
            bottom=Side(style='thin')
        )
        
        # Tiêu đề chính
        ws.merge_cells('A1:B1')
        ws['A1'] = 'DECISION TREE MODEL - EVALUATION METRICS'
        ws['A1'].font = Font(bold=True, size=14)
        ws['A1'].alignment = cell_alignment
        ws['A1'].fill = PatternFill(start_color="385723", end_color="385723", fill_type="solid")
        ws['A1'].font = Font(bold=True, color="FFFFFF", size=14)
        
        # Header cho metrics tổng quan
        ws['A3'] = 'Metric'
        ws['B3'] = 'Value'
        ws['A3'].fill = header_fill
        ws['B3'].fill = header_fill
        ws['A3'].font = header_font
        ws['B3'].font = header_font
        ws['A3'].alignment = cell_alignment
        ws['B3'].alignment = cell_alignment
        
        # Dữ liệu metrics
        metrics_data = [
            ['Accuracy', f"{results['test_accuracy']:.4f}"],
            ['Precision', f"{results['precision']:.4f}"],
            ['Recall', f"{results['recall']:.4f}"],
            ['F1-Score', f"{results['f1_score']:.4f}"],
            ['AUC', f"{results['roc_auc']:.4f}"]
        ]
        
        row = 4
        for metric_name, metric_value in metrics_data:
            ws[f'A{row}'] = metric_name
            ws[f'B{row}'] = metric_value
            ws[f'A{row}'].alignment = cell_alignment
            ws[f'B{row}'].alignment = cell_alignment
            ws[f'A{row}'].border = border
            ws[f'B{row}'].border = border
            row += 1
        
        # Metrics per class
        row += 2
        ws.merge_cells(f'A{row}:D{row}')
        ws[f'A{row}'] = 'METRICS PER CLASS'
        ws[f'A{row}'].font = Font(bold=True, size=12)
        ws[f'A{row}'].alignment = cell_alignment
        ws[f'A{row}'].fill = PatternFill(start_color="385723", end_color="385723", fill_type="solid")
        ws[f'A{row}'].font = Font(bold=True, color="FFFFFF", size=12)
        
        row += 1
        headers = ['Class', 'Precision', 'Recall', 'F1-Score']
        for col, header in enumerate(headers, start=1):
            cell = ws.cell(row=row, column=col)
            cell.value = header
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = cell_alignment
            cell.border = border
        
        row += 1
        class_names = ['Class 0 (Không bệnh)', 'Class 1 (Có bệnh)']
        for i, class_name in enumerate(class_names):
            ws.cell(row=row, column=1, value=class_name)
            ws.cell(row=row, column=2, value=f"{results['precision_per_class'][i]:.4f}")
            ws.cell(row=row, column=3, value=f"{results['recall_per_class'][i]:.4f}")
            ws.cell(row=row, column=4, value=f"{results['f1_per_class'][i]:.4f}")
            
            for col in range(1, 5):
                ws.cell(row=row, column=col).alignment = cell_alignment
                ws.cell(row=row, column=col).border = border
            row += 1
        
        # Cross-Validation Scores
        row += 2
        ws.merge_cells(f'A{row}:B{row}')
        ws[f'A{row}'] = 'CROSS-VALIDATION SCORES'
        ws[f'A{row}'].font = Font(bold=True, size=12)
        ws[f'A{row}'].alignment = cell_alignment
        ws[f'A{row}'].fill = PatternFill(start_color="385723", end_color="385723", fill_type="solid")
        ws[f'A{row}'].font = Font(bold=True, color="FFFFFF", size=12)
        
        row += 1
        ws[f'A{row}'] = 'Fold'
        ws[f'B{row}'] = 'Score'
        ws[f'A{row}'].fill = header_fill
        ws[f'B{row}'].fill = header_fill
        ws[f'A{row}'].font = header_font
        ws[f'B{row}'].font = header_font
        ws[f'A{row}'].alignment = cell_alignment
        ws[f'B{row}'].alignment = cell_alignment
        
        row += 1
        for i, score in enumerate(results['cv_scores'], 1):
            ws[f'A{row}'] = f'Fold {i}'
            ws[f'B{row}'] = f'{score:.4f}'
            ws[f'A{row}'].alignment = cell_alignment
            ws[f'B{row}'].alignment = cell_alignment
            ws[f'A{row}'].border = border
            ws[f'B{row}'].border = border
            row += 1
        
        ws[f'A{row}'] = 'Mean'
        ws[f'B{row}'] = f'{results["cv_scores"].mean():.4f}'
        ws[f'A{row}'].font = Font(bold=True)
        ws[f'B{row}'].font = Font(bold=True)
        ws[f'A{row}'].alignment = cell_alignment
        ws[f'B{row}'].alignment = cell_alignment
        ws[f'A{row}'].border = border
        ws[f'B{row}'].border = border
        
        row += 1
        ws[f'A{row}'] = 'Std Dev'
        ws[f'B{row}'] = f'{results["cv_scores"].std():.4f}'
        ws[f'A{row}'].font = Font(bold=True)
        ws[f'B{row}'].font = Font(bold=True)
        ws[f'A{row}'].alignment = cell_alignment
        ws[f'B{row}'].alignment = cell_alignment
        ws[f'A{row}'].border = border
        ws[f'B{row}'].border = border
        
        # Confusion Matrix
        row += 3
        ws.merge_cells(f'A{row}:C{row}')
        ws[f'A{row}'] = 'CONFUSION MATRIX'
        ws[f'A{row}'].font = Font(bold=True, size=12)
        ws[f'A{row}'].alignment = cell_alignment
        ws[f'A{row}'].fill = PatternFill(start_color="385723", end_color="385723", fill_type="solid")
        ws[f'A{row}'].font = Font(bold=True, color="FFFFFF", size=12)
        
        row += 1
        ws[f'A{row}'] = ''
        ws[f'B{row}'] = 'Predicted 0'
        ws[f'C{row}'] = 'Predicted 1'
        for col in ['A', 'B', 'C']:
            ws[f'{col}{row}'].fill = header_fill
            ws[f'{col}{row}'].font = header_font
            ws[f'{col}{row}'].alignment = cell_alignment
            ws[f'{col}{row}'].border = border
        
        cm = results['confusion_matrix']
        row += 1
        ws[f'A{row}'] = 'Actual 0'
        ws[f'B{row}'] = cm[0][0]
        ws[f'C{row}'] = cm[0][1]
        for col in ['A', 'B', 'C']:
            ws[f'{col}{row}'].alignment = cell_alignment
            ws[f'{col}{row}'].border = border
        ws[f'A{row}'].font = Font(bold=True)
        
        row += 1
        ws[f'A{row}'] = 'Actual 1'
        ws[f'B{row}'] = cm[1][0]
        ws[f'C{row}'] = cm[1][1]
        for col in ['A', 'B', 'C']:
            ws[f'{col}{row}'].alignment = cell_alignment
            ws[f'{col}{row}'].border = border
        ws[f'A{row}'].font = Font(bold=True)
        
        # Feature Importance
        row += 3
        ws.merge_cells(f'A{row}:B{row}')
        ws[f'A{row}'] = 'FEATURE IMPORTANCE (Top 10)'
        ws[f'A{row}'].font = Font(bold=True, size=12)
        ws[f'A{row}'].alignment = cell_alignment
        ws[f'A{row}'].fill = PatternFill(start_color="385723", end_color="385723", fill_type="solid")
        ws[f'A{row}'].font = Font(bold=True, color="FFFFFF", size=12)
        
        row += 1
        ws[f'A{row}'] = 'Feature'
        ws[f'B{row}'] = 'Importance'
        ws[f'A{row}'].fill = header_fill
        ws[f'B{row}'].fill = header_fill
        ws[f'A{row}'].font = header_font
        ws[f'B{row}'].font = header_font
        ws[f'A{row}'].alignment = cell_alignment
        ws[f'B{row}'].alignment = cell_alignment
        
        row += 1
        for _, feat_row in results['feature_importance'].head(10).iterrows():
            ws[f'A{row}'] = feat_row['feature']
            ws[f'B{row}'] = f"{feat_row['importance']:.4f}"
            ws[f'A{row}'].alignment = cell_alignment
            ws[f'B{row}'].alignment = cell_alignment
            ws[f'A{row}'].border = border
            ws[f'B{row}'].border = border
            row += 1
        
        # Điều chỉnh độ rộng cột
        ws.column_dimensions['A'].width = 25
        ws.column_dimensions['B'].width = 15
        ws.column_dimensions['C'].width = 15
        ws.column_dimensions['D'].width = 15
        
        # Lưu file
        wb.save(filename)
        print(f"✓ Đã xuất metrics ra file: {filename}")
        
        return filename
